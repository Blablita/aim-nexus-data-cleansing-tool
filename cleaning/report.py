"""Builds the Excel data quality report.

Sheets always present:  Summary, Detailed_Log
Optional sheets (when the config turns those steps on):
    Control_Totals, Investor_Summary, Top_Investors, Cross_File,
    EDA_Country, EDA_Entity, EDA_Home_vs_Intl, EDA_Holding_Size, EDA_Statistics
"""

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BAD_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
OK_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _style(ws):
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        cells = list(column_cells)[:200]
        length = max(len(str(c.value)) if c.value is not None else 0 for c in cells)
        ws.column_dimensions[cells[0].column_letter].width = min(max(length + 2, 10), 60)
    # color reconciliation results (only on small sheets: on 30,000-row sheets this is slow
    # and the colors add nothing)
    if ws.max_row <= 500:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value in ("MISMATCH", "CHECK"):
                    c.fill = BAD_FILL
                elif c.value == "OK":
                    c.fill = OK_FILL


def build_excel_report(summary_rows, detail_rows, output_path, extras=None, top_n=200):
    """
    summary_rows: one row per file.  detail_rows: one event per row.
    extras: optional dict with per-file lists:
        reconciliation, investors (DataFrames), collation_info (dicts), eda (dicts of DataFrames),
        cross_file (DataFrame)
    """
    extras = extras or {}
    sheets = {"Summary": pd.DataFrame(summary_rows),
              "Detailed_Log": pd.DataFrame(detail_rows) if detail_rows else pd.DataFrame([{"info": "No events."}])}
    if extras.get("reconciliation"):
        sheets["Control_Totals"] = pd.concat(extras["reconciliation"], ignore_index=True)
    if extras.get("collation_info"):
        sheets["Investor_Summary"] = pd.DataFrame(extras["collation_info"])
    if extras.get("investors"):
        sheets["Top_Investors"] = pd.concat([t.head(top_n) for t in extras["investors"]], ignore_index=True)
    if extras.get("cross_file") is not None and len(extras["cross_file"]):
        sheets["Cross_File"] = extras["cross_file"]
    for key, name in [("by_country", "EDA_Country"), ("by_entity", "EDA_Entity"),
                      ("home_vs_international", "EDA_Home_vs_Intl"), ("holding_size", "EDA_Holding_Size"),
                      ("statistics", "EDA_Statistics")]:
        parts = [e[key] for e in extras.get("eda", []) if key in e]
        if parts:
            sheets[name] = pd.concat(parts, ignore_index=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            _style(writer.sheets[name])
    return output_path
